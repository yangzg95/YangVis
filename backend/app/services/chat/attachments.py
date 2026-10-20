"""对话附件的落盘、解析与回放读取。

文档不存原文：上传时就抽取成纯文本，后续每一轮都从这里取。图片存原文件，
经 ``/uploads`` 静态路径给前端预览，回放时读回来转成 data URL 交给视觉模型。

落盘规则与幻灯片配图一致（见 ``services/slides/store.py``）：扩展名按内容
魔数判定、文件名由服务端随机生成、删除前按记录里的相对路径复验。用户给的
文件名只用于展示，永远不参与路径拼接。
"""
from __future__ import annotations

import base64
import io
import logging
import re
import uuid
from dataclasses import dataclass

from app.config import Settings, get_settings
from app.services.chunking import normalise
from app.services.slides.store import SlideImageError, sniff_image

logger = logging.getLogger("yangvis.chat.attachments")

settings: Settings = get_settings()

# 能直接抽出文本的格式。PDF / DOCX 走专用解析器，其余按文本解码。
TEXT_EXTENSIONS = (
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".log",
    ".yaml", ".yml", ".xml", ".toml", ".ini", ".conf", ".properties", ".env",
    ".sql", ".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".vue",
    ".java", ".kt", ".go", ".rs", ".c", ".h", ".cpp", ".hpp", ".cs",
    ".php", ".rb", ".sh", ".bash", ".zsh", ".bat", ".ps1",
    ".html", ".htm", ".css", ".scss", ".less",
)
SUPPORTED_FILE_EXTENSIONS = (".pdf", ".docx") + TEXT_EXTENSIONS

# 至少保留一份可读的说明，别让用户只看到一串白名单。
SUPPORTED_FILE_HINT = "PDF / Word（.docx）/ 纯文本 / 常见代码文件"

# chat/<owner id>/<32 位十六进制>.<真实扩展名>。写入方是本模块，但库里的
# 值可能被任何途径改过，读取前仍然复验一次。
_REL_PATH_RE = re.compile(
    r"^chat/(?P<owner>\d{1,20})/(?P<name>[0-9a-f]{32})\.(?P<ext>png|jpg|gif|webp)$"
)


class AttachmentError(ValueError):
    """附件不可用（格式不支持、解析失败、写入失败）。路由按业务拒绝呈现。"""


@dataclass(frozen=True)
class StoredChatImage:
    mime: str
    size: int
    rel_path: str
    url: str


def _truncate(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return f"{text[:limit]}\n……（内容过长，已截断）"


def _decode_text(filename: str, raw: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "utf-16"):
        try:
            return normalise(raw.decode(encoding))
        except UnicodeDecodeError:
            continue
    logger.debug("could not detect the encoding of %s, decoding with replacement", filename)
    return normalise(raw.decode("utf-8", errors="replace"))


def extract_file_text(filename: str, raw: bytes) -> str:
    """把文档抽取成文本；图片不走这里。

    返回的文本已按 ``CHAT_ATTACHMENT_MAX_CHARS`` 截断。空文本视为解析失败：
    一张扫描版 PDF 存进去只会让模型收到一条空附件，不如当场告诉用户。
    """
    lowered = filename.lower()

    if lowered.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(raw))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:  # noqa: BLE001 - pypdf 的异常类型随文件而坏法各异
            raise AttachmentError(f"PDF 解析失败：{exc}") from exc
        text = normalise(text)
        if not text:
            raise AttachmentError("PDF 中没有可提取的文本，扫描件请转成图片上传")
        return _truncate(text, settings.CHAT_ATTACHMENT_MAX_CHARS)

    if lowered.endswith(".docx"):
        import docx

        try:
            document = docx.Document(io.BytesIO(raw))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        except Exception as exc:  # noqa: BLE001 - python-docx 对坏文件抛的异常很杂
            raise AttachmentError(f"Word 文档解析失败：{exc}") from exc
        text = normalise("\n".join(parts))
        if not text:
            raise AttachmentError("Word 文档中没有可提取的文本")
        return _truncate(text, settings.CHAT_ATTACHMENT_MAX_CHARS)

    if lowered.endswith(TEXT_EXTENSIONS):
        text = _decode_text(filename, raw)
        if not text:
            raise AttachmentError("文件内容为空")
        return _truncate(text, settings.CHAT_ATTACHMENT_MAX_CHARS)

    raise AttachmentError(
        f"暂不支持该文件格式，目前支持 {SUPPORTED_FILE_HINT}"
    )


def store_image(owner_id: int, filename: str, raw: bytes) -> StoredChatImage:
    """把图片写进 ``chat/<owner>/<随机 id>.<真实扩展名>``。

    只收视觉模型普遍支持的格式。BMP 虽然在滑片素材里允许，但多数对话模型
    不接受，这里直接拒绝，而不是等上游整条请求失败。
    """
    if len(raw) > settings.CHAT_MAX_IMAGE_BYTES:
        raise AttachmentError(
            f"图片超过 {settings.CHAT_MAX_IMAGE_BYTES // 1024 // 1024}MB 上限"
        )
    try:
        mime, extension = sniff_image(raw)
    except SlideImageError as exc:
        raise AttachmentError(str(exc)) from exc
    if mime == "image/bmp":
        raise AttachmentError("图片请使用 PNG / JPEG / GIF / WebP 格式")

    relative = f"chat/{owner_id}/{uuid.uuid4().hex}{extension}"
    target = settings.upload_media_path / relative
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(raw)
    except OSError as exc:
        logger.error("chat image write failed: %s (%s)", target, exc)
        raise AttachmentError(f"图片写入服务器失败：{exc}") from exc

    logger.info(
        "chat image stored: owner=%s %s %s bytes=%d",
        owner_id,
        relative,
        mime,
        len(raw),
    )
    return StoredChatImage(
        mime=mime,
        size=len(raw),
        rel_path=relative,
        url=image_url(relative),
    )


def is_safe_rel_path(rel_path: str) -> bool:
    return bool(_REL_PATH_RE.match(rel_path))


def image_url(rel_path: str) -> str:
    """由库里的相对路径拼浏览器 URL；不合法时返回空串。"""
    if not is_safe_rel_path(rel_path):
        logger.warning("chat image url rejected, unsafe rel_path: %r", rel_path)
        return ""
    return f"{settings.UPLOAD_MEDIA_URL.rstrip('/')}/{rel_path}"


def read_image_data_url(rel_path: str, mime: str | None) -> str | None:
    """读回图片并转成 data URL；文件丢失或路径不合法时返回 None。"""
    if not is_safe_rel_path(rel_path):
        logger.warning("chat image replay rejected, unsafe rel_path: %r", rel_path)
        return None
    path = settings.upload_media_path / rel_path
    try:
        raw = path.read_bytes()
    except OSError as exc:
        logger.warning("chat image replay failed: %s (%s)", path, exc)
        return None
    encoded = base64.b64encode(raw).decode("ascii")
    return f"data:{mime or 'image/png'};base64,{encoded}"


def delete_stored(rel_path: str | None) -> None:
    """best-effort 删掉图片文件；数据库行已经删了，磁盘残留不阻塞流程。"""
    if not rel_path or not is_safe_rel_path(rel_path):
        if rel_path:
            logger.warning("chat attachment delete skipped, unsafe rel_path: %r", rel_path)
        return
    path = (settings.upload_media_path / rel_path).resolve()
    try:
        path.unlink(missing_ok=True)
    except OSError as exc:
        logger.warning("chat attachment delete failed: %s (%s)", path, exc)
